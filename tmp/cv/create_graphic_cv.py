from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

assets=Path('tmp/cv/graphic_assets'); assets.mkdir(parents=True,exist_ok=True)
navy='16364A'; teal='007F87'; pale='EAF4F4'
# Decorative geometric artwork and monochrome pictograms.
im=Image.new('RGB',(420,420),'#'+navy); d=ImageDraw.Draw(im)
d.ellipse((50,-90,520,380),fill='#'+teal);d.ellipse((145,5,425,285),fill='#'+navy)
d.polygon([(0,290),(210,80),(210,180),(0,390)],fill='#62D3C7')
im.save(assets/'geometry.png')
for name in ['pin','phone','mail','code','work','school','tools','language']:
 im=Image.new('RGBA',(100,100),(0,0,0,0));d=ImageDraw.Draw(im);c='#'+teal
 if name=='mail':
  d.rounded_rectangle((14,25,86,76),radius=6,outline=c,width=6);d.line([(16,28),(50,54),(84,28)],fill=c,width=6)
 elif name=='pin':
  d.ellipse((25,10,75,60),outline=c,width=6);d.line([(27,48),(50,90),(73,48)],fill=c,width=6);d.ellipse((42,27,58,43),fill=c)
 elif name=='phone':
  d.rounded_rectangle((29,8,71,92),radius=7,outline=c,width=6);d.line((42,80,58,80),fill=c,width=5)
 elif name=='code':
  d.line([(33,23),(12,50),(33,77)],fill=c,width=6);d.line([(67,23),(88,50),(67,77)],fill=c,width=6);d.line((57,16,43,84),fill=c,width=6)
 elif name=='work':
  d.rounded_rectangle((10,32,90,85),radius=6,outline=c,width=6);d.rectangle((34,16,66,33),outline=c,width=6);d.line((12,52,88,52),fill=c,width=5)
 elif name=='school':
  d.polygon([(8,37),(50,15),(92,37),(50,59)],outline=c,width=6);d.line([(25,49),(25,72),(50,85),(75,72),(75,49)],fill=c,width=6)
 elif name=='tools':
  d.line((22,80,77,25),fill=c,width=13);d.arc((48,8,92,52),0,285,fill=c,width=7)
 else:
  d.rounded_rectangle((10,15,90,73),radius=10,outline=c,width=6);d.line([(28,73),(25,90),(48,73)],fill=c,width=6);d.line((27,35,73,35),fill=c,width=5);d.line((27,52,60,52),fill=c,width=5)
 im.save(assets/f'{name}.png')
for count in [1,2]:
 im=Image.new('RGB',(500,30),'#'+pale);d=ImageDraw.Draw(im)
 for i in range(2):d.rounded_rectangle((i*255,0,i*255+245,29),radius=10,fill='#'+(teal if i<count else 'C6DDDE'))
 im.save(assets/f'bar{count}.png')

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(.9);sec.bottom_margin=Cm(.9);sec.left_margin=Cm(1);sec.right_margin=Cm(1)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9.5)
normal.font.color.rgb=RGBColor.from_string('243746')
normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.08
for style in ['Title','Heading 1','Heading 2']:
 doc.styles[style].font.name='Arial';doc.styles[style].font.color.rgb=RGBColor(0,0,0)
for el in [doc.styles.element,doc.element]:
 for b in list(el.iter(qn('w:pBdr'))):b.getparent().remove(b)

def shade(cell,color):
 x=OxmlElement('w:shd');x.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(x)
def margins(cell,top=180,start=200,bottom=180,end=200):
 x=OxmlElement('w:tcMar')
 for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');x.append(e)
 cell._tc.get_or_add_tcPr().append(x)
def table_clean(table,widths):
 table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+k);x.set(qn('w:val'),'nil');borders.append(x)
 table._tbl.tblPr.append(borders)
 for col,w in zip(table.columns,widths):col.width=Cm(w)
 for row in table.rows:
  for cell,w in zip(row.cells,widths):cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

def para(cell,text='',size=9.5,bold=False,color='243746',after=4,before=0,style=None):
 p=cell.add_paragraph(style=style);p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(before)
 r=p.add_run(text);r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 return p

def heading(cell,text,icon):
 p=cell.add_paragraph();p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(assets/f'{icon}.png'),width=Cm(.38))
 r=p.add_run('  '+text);r.font.size=Pt(11);r.bold=True;r.font.color.rgb=RGBColor(0,0,0)

def empty_remove(cell):
 p=cell.paragraphs[0];p._element.getparent().remove(p._element)

banner=doc.add_table(rows=1,cols=2);table_clean(banner,[15.5,3.5])
for cell in banner.rows[0].cells:shade(cell,navy);margins(cell,200,300,200,300);empty_remove(cell)
a,b=banner.rows[0].cells
para(a,'Modou Mamoune Diop',25,True,'FFFFFF',6,style='Title')
para(a,'Développeur Flutter et web',13,False,'72DDD0',0)
p=b.add_paragraph();p.add_run().add_picture(str(assets/'geometry.png'),width=Cm(2.5))
gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(6);gap.add_run().font.size=Pt(1)
body=doc.add_table(rows=1,cols=2);table_clean(body,[6,13]);left,right=body.rows[0].cells
for cell in [left,right]:margins(cell,100,250,140,250);empty_remove(cell)
shade(left,pale)
heading(left,'Contact','mail')
for icon,text in [('pin','Dakar, Sénégal'),('phone','+221 77 157 93 13'),('mail','mamoune044@gmail.com'),('code','github.com/DiopMamoune')]:
 p=left.add_paragraph();p.paragraph_format.space_after=Pt(6)
 p.add_run().add_picture(str(assets/f'{icon}.png'),width=Cm(.28));r=p.add_run('  '+text);r.font.size=Pt(8.5)
para(left,'LinkedIn',9,True,after=2)
para(left,'linkedin.com/in/\nmodou-mamoune-diop-598952199',8,after=4)
heading(left,'Compétences','code')
para(left,'Missions par technologie\nÉchelle : 0 à 2 missions',8,color='4A6570',after=7)
for tech,count in [('Flutter / Dart',2),('Angular',2),('React / Next.js',1),('PHP',1),('Spring Boot',1)]:
 para(left,f'{tech}   |   {count} mission'+('s' if count>1 else ''),9,True,after=2)
 p=left.add_paragraph();p.paragraph_format.space_after=Pt(7);p.add_run().add_picture(str(assets/f'bar{count}.png'),width=Cm(4.5),height=Cm(.12))
heading(left,'Outils','tools')
para(left,'Git · Bitbucket · GitLab\nGitHub · Jira\nPostman · Insomnia',9,after=4)
heading(left,'Formation','school')
for date,degree,school in [('2020 - 2021','Certification Développement\nWeb et Mobile','Access Code School'),('2016 - 2019','Licence Génie Logiciel','ISI'),('2015 - 2016','Baccalauréat série S2','Lycée Seydina Limamou Laye')]:
 para(left,date,8,color=teal,after=2)
 para(left,degree,9,True,after=2)
 para(left,school,8.5,after=8)
heading(left,'Langues','language')
para(left,'Wolof : langue maternelle\nFrançais : intermédiaire\nAnglais : intermédiaire',9)

heading(right,'Profil','code')
para(right,'Je développe des applications mobiles avec Flutter et Dart et des interfaces web avec React, Next.js et Angular. Mes missions chez CBAO et Finetech portent sur le paiement interopérable et la gestion des dépenses publiques.',9.5,after=7)
heading(right,'Expérience professionnelle','work')
jobs=[
('CBAO','Groupe Attijariwafa Bank','Consultant développeur mobile Flutter','Septembre 2024 - Présent | Dakar, sur site',['Développement d’une application mobile intégrant le système de paiement interopérable de la BCEAO pour la zone UEMOA.'],'Flutter · Dart · Bitbucket · Jira · Agile'),
('Finetech Groupe','','Développeur web et mobile','Février 2023 - Présent',['Travail sur le module d’exécution des dépenses et de comptabilité budgétaire de Sysbudgep pour la Direction Générale du Budget du Sénégal.'],'React · Next.js · Bitbucket'),
('GITECH SN','','Développeur full-stack','Août 2022 - Janvier 2023',['Développement d’une plateforme de gestion de la mutuelle de santé des militaires sénégalais.'],'Angular 12 · PHP 8 · GitLab'),
('FINDME','','Développeur full-stack','Juillet 2022 - Décembre 2022',['Développement d’une plateforme de gestion et de demandes d’agréments.','Développement d’une application mobile générant une adresse postale géolocalisée à partir de la position GPS.'],'Flutter · Sentry'),
('UNITED VENDOR','','Développeur full-stack','Avril 2021 - Décembre 2021',['Mise en place de Mifos, solution open source de gestion bancaire, et personnalisation selon les besoins clients.'],'Angular · Spring Boot')]
for company,group,role,dates,items,stack in jobs:
 p=para(right,company,12,True,after=3,before=8);p.paragraph_format.keep_with_next=True
 if group:para(right,group,9,color='4A6570',after=3)
 para(right,role,9.5,True,after=3)
 para(right,dates,8.5,color=teal,after=6)
 for text in items:
  p=para(right,'• '+text,9.5,after=4)
  p.paragraph_format.left_indent=Cm(.2);p.paragraph_format.first_line_indent=Cm(-.2)
 para(right,stack,8.5,False,teal,after=6)

doc.core_properties.author='Modou Mamoune Diop';doc.core_properties.title='CV graphique Modou Mamoune Diop'
doc.core_properties.subject='Développeur Flutter et web'
doc.save('output/cv/CV_Modou_Mamoune_Diop_Graphique.docx')
