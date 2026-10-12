from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('output/cv')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.4); sec.bottom_margin=Cm(1.4)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
normal=doc.styles['Normal']
normal.font.name='Arial'; normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(3)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Subtitle',13),('Heading 1',11),('Heading 2',10)]:
    st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(10 if name=='Heading 1' else 3)
    st.paragraph_format.space_after=Pt(4)
    if name!='Subtitle': st.font.bold=True
    st.font.italic=False
    st.paragraph_format.keep_with_next=True

def p(text,bold=False,after=3):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(after)
    para.add_run(text).bold=bold
    return para

def bullet(text):
    para=doc.add_paragraph(text,style='List Bullet')
    para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.left_indent=Cm(.35)
    para.paragraph_format.first_line_indent=Cm(-.2)

def job(role,company,dates,items,stack):
    para=p(company,True,2)
    para.paragraph_format.keep_with_next=True
    para.runs[0].font.size=Pt(11)
    para=p(role+' | '+dates,False,4)
    para.paragraph_format.keep_with_next=True
    para.runs[0].font.size=Pt(9)
    para.runs[0].font.color.rgb=RGBColor.from_string('344054')
    for item in items: bullet(item)
    para=p(stack,False,7)
    para.runs[0].font.size=Pt(9)
    para.runs[0].font.color.rgb=RGBColor.from_string('183B56')

# All contact text is in the body so it remains extractable by recruitment systems.
doc.add_paragraph('Modou Mamoune Diop','Title')
subtitle=doc.add_paragraph('Développeur Flutter et web','Subtitle')
subtitle.runs[0].font.color.rgb=RGBColor.from_string('183B56')
p('Dakar, Sénégal | +221 77 157 93 13 | mamoune044@gmail.com',after=2)
p('github.com/DiopMamoune | linkedin.com/in/modou-mamoune-diop-598952199',after=5).runs[0].font.size=Pt(9)

doc.add_heading('Profil',1)
p('Je développe des applications mobiles avec Flutter et Dart et des interfaces web avec React, Next.js et Angular. Mes missions chez CBAO et Finetech portent sur le paiement interopérable et la gestion des dépenses publiques. J’ai également travaillé sur une mutuelle de santé et des services de géolocalisation.')

doc.add_heading('Expérience professionnelle',1)
job('Consultant développeur mobile Flutter','CBAO - Groupe Attijariwafa Bank','Septembre 2024 - Présent | Dakar, sur site',[
    'Développement d’une application mobile intégrant le système de paiement interopérable de la BCEAO pour la zone UEMOA.'
],'Flutter, Dart, Bitbucket, Jira, Agile')
job('Développeur web et mobile','Finetech Groupe','Février 2023 - Présent',[
    'Travail sur le module d’exécution des dépenses et de comptabilité budgétaire de Sysbudgep pour la Direction Générale du Budget du Sénégal.'
],'React, Next.js, Bitbucket')
job('Développeur full-stack','GITECH SN','Août 2022 - Janvier 2023',[
    'Développement d’une plateforme de gestion de la mutuelle de santé des militaires sénégalais.'
],'Angular 12, PHP 8, GitLab')
job('Développeur full-stack','FINDME','Juillet 2022 - Décembre 2022',[
    'Développement d’une plateforme de gestion et de demandes d’agréments.',
    'Développement d’une application mobile générant une adresse postale géolocalisée à partir de la position GPS.'
],'Flutter, Sentry')
job('Développeur full-stack','UNITED VENDOR','Avril 2021 - Décembre 2021',[
    'Mise en place de Mifos, solution open source de gestion bancaire, et personnalisation selon les besoins clients.'
],'Angular, Spring Boot')

doc.add_heading('Compétences techniques',1)
p('Mobile : Flutter, Dart | Web : React, Next.js, Angular | Backend : PHP, Spring Boot')
p('Outils : Git, Bitbucket, GitLab, GitHub, Jira, Postman, Insomnia')

doc.add_heading('Formation',1)
p('2020 - 2021 | Certification Développement Web et Mobile | Access Code School')
p('2016 - 2019 | Licence Génie Logiciel | ISI')
p('2015 - 2016 | Baccalauréat série S2 | Lycée Seydina Limamou Laye')

doc.add_heading('Langues',1)
p('Wolof : langue maternelle | Français : intermédiaire | Anglais : intermédiaire')

doc.core_properties.author='Modou Mamoune Diop'
doc.core_properties.title='CV Modou Mamoune Diop'
doc.core_properties.subject='Développeur web et mobile'
doc.core_properties.keywords='Flutter, Dart, React, Next.js, Angular, PHP, Spring Boot'
for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(out/'CV_Modou_Mamoune_Diop_Moderne.docx')
